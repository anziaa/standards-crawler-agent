"""把批量爬取结果汇总回编制依据表，输出带「爬取结果/失败原因/文件位置/校验结果」的新 Excel。

用法：
    set PYTHONPATH=src
    python -m standards_crawler_agent.report            # 汇总 results/batch_sheet_* 全部批次
    python -m standards_crawler_agent.report --results results/batch_sheet_001 --output out.xlsx

设计要点：
  1. **按名称匹配，不按序号**：batch 驱动会去掉重名条目，任务目录里的序号会因此错位；
     名称（去掉空白后）才是稳定键。表里 13 组完全重名的行会共用同一份结果。
  2. **保留原表**：读入原工作簿后只往目标工作表追加列，其余工作表原样保留，
     另存为新文件，绝不覆盖原文件。
  3. **人工可读优先**：状态列写中文判断，同时保留英文原值列便于追溯；
     另附「问题清单」工作表，把需要人看的行挑出来。
"""

from __future__ import annotations

import argparse
import json
import os
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 编制依据表是各人自己的业务文件，不进版本库，所以默认值只写**相对路径**（相对当前工作目录）。
# 用环境变量 SHEET_SOURCE / SHEET_NAME 覆盖（推荐写进 .env），或用命令行 --source / --sheet 覆盖。
DEFAULT_SOURCE = Path(os.getenv("SHEET_SOURCE", "编制依据汇总表.xlsx"))
SHEET_NAME = os.getenv("SHEET_NAME", "编制依据")
DEFAULT_RESULT_GLOB = "results/batch_sheet_*"

# 追加列（原表三列之后）
COLUMNS: list[tuple[str, int]] = [
    ("行号", 6),
    ("爬取结果", 14),
    ("命中轮次", 12),
    ("失败原因", 42),
    ("失败归类", 24),
    ("文件位置", 30),
    ("文件名", 30),
    ("校验结果", 30),
    ("校验提示", 32),
    ("过程提示", 34),
    ("识别到的文号", 20),
    ("内容哈希(sha256)", 20),
    ("类型", 10),
    ("页数", 8),
    ("大小(KB)", 10),
    ("来源页面", 40),
    ("耗时(秒)", 10),
    ("token", 9),
    ("需人工复核", 12),
    ("人工结论", 14),
    ("结果状态(原始)", 20),
]

STATUS_LABEL: dict[str, str] = {
    "success": "成功（官方原件）",
    "success_page_export": "成功（网页导出）",
    "success_scanned": "成功（扫描件，无文本层）",
    "not_downloadable": "未找到可下载文件",
    "needs_confirmation": "需人工确认",
    "skipped_withdrawn": "已废止/失效（未下载）",
    "timeout": "超时",
    "error": "异常",
}

STATUS_FILL: dict[str, str] = {
    "success": "C6EFCE",  # 绿
    "success_page_export": "FFEB9C",  # 黄：可用但不是官方原件
    "success_scanned": "DDEBF7",  # 蓝：内容在，但没有文本层可比对
    "needs_confirmation": "FFEB9C",
    "skipped_withdrawn": "D9D9D9",  # 灰
    "not_downloadable": "FFC7CE",  # 红
    "timeout": "FFC7CE",
    "error": "FFC7CE",
}


def normalize(value: str | None) -> str:
    return re.sub(r"[\s\u3000]+", "", value or "").lower()


def load_run(results_roots: list[Path]) -> tuple[dict[str, dict[str, Any]], list[str]]:
    """读取若干批次目录，返回 {规范化名称: 结果记录} 与运行顺序说明。"""
    records: dict[str, dict[str, Any]] = {}
    runs: list[str] = []
    for root in results_roots:
        summary_path = root / "batch_summary.json"
        if not summary_path.exists():
            continue
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
        runs.append(f"{root}（墙钟 {summary.get('total_seconds')}s，token {summary.get('total_tokens')}，任务 {len(summary.get('results', []))}）")
        # 目录名里带 retry 的批次算"重跑"，用于区分一次命中与重跑命中。
        run_label = "重跑" if "retry" in root.name else "首次"
        for item in summary.get("results", []):
            key = normalize(item.get("filename"))
            # 同一名称出现在多个批次时，后跑的覆盖先跑的（便于重爬单条后重新汇总）。
            records[key] = {**item, "run_root": str(root), "run_label": run_label}
    return records, runs


def read_task_detail(record: dict[str, Any]) -> dict[str, Any]:
    """读任务目录里的 result.json，取 verification / trace 等明细。"""
    task_dir = record.get("task_dir")
    if not task_dir:
        return {}
    # summary 里存的是相对项目根目录的路径；这里按绝对路径与相对路径都试一次。
    candidates = [Path(task_dir)]
    if not Path(task_dir).is_absolute():
        candidates.append(Path.cwd() / task_dir)
    for path in candidates:
        result_file = path / "result.json"
        if result_file.exists():
            try:
                return json.loads(result_file.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                return {}
    return {}


def summarize(record: dict[str, Any], detail: dict[str, Any]) -> dict[str, Any]:
    status = str(record.get("status") or "")
    verification = detail.get("verification") or {}
    metadata = verification.get("metadata") or {}
    warnings = list(verification.get("warnings") or [])
    notes = list(verification.get("notes") or [])
    errors = [str(item) for item in (record.get("errors") or detail.get("errors") or [])]
    skip_reason = detail.get("skip_reason") or ""
    quarantined = metadata.get("quarantined_to")

    path = record.get("downloaded_path") or detail.get("downloaded_path")
    location = ""
    if path:
        location = str(path)
    if quarantined:
        location = f"（已拒绝）{quarantined}"

    checks: list[str] = []
    if verification:
        checks.append("内容校验通过" if verification.get("ok") else "内容校验未通过")
        if metadata.get("name_match"):
            checks.append(str(metadata["name_match"]))
        if metadata.get("pages"):
            checks.append(f"{metadata['pages']} 页")
    elif status in {"timeout", "error"}:
        checks.append("未执行校验（任务未产出文件）")
    elif status == "not_downloadable":
        checks.append("未下载，无需校验")
    elif status == "skipped_withdrawn":
        checks.append("已判废止，未下载")

    needs_review, review_reason = review_flag(status, warnings, notes, errors, skip_reason, bool(quarantined), bool(verification))

    # 成功的行不该把过程性提示写进「失败原因」——那会让人误以为出了问题。
    # 检索被丢弃、生命周期判定说明这类信息归到「过程提示」。
    failing = status in {"not_downloadable", "needs_confirmation", "timeout", "error", "skipped_withdrawn"}
    reason_parts: list[str] = []
    note_parts: list[str] = []
    for item in errors:
        (reason_parts if failing else note_parts).append(item)
    if skip_reason:
        (reason_parts if failing else note_parts).append(skip_reason)

    return {
        "label": STATUS_LABEL.get(status, status or "未知"),
        "run_label": record.get("run_label", ""),
        "status": status,
        "location": location,
        "filename": Path(str(path)).name if path else "",
        "check": "；".join(checks),
        # 非致命的观察（扫描件无文本层等）与致命警告写在同一列，但保留「提示」前缀，
        # 让人一眼看出它不是校验失败。
        "warnings": "；".join(warnings + [f"提示：{item}" for item in notes]),
        "reason": "；".join(reason_parts),
        "notes": "；".join(note_parts),
        "failure_kind": classify_failure(status, errors, reason_parts),
        "standard_number": metadata.get("standard_number", ""),
        "sha256": metadata.get("sha256", ""),
        "mime": verification.get("mime_type", ""),
        "pages": metadata.get("pages", record.get("verification_pages") or ""),
        "size": verification.get("size") or record.get("verification_size") or 0,
        "source": record.get("source_url") or detail.get("source_url") or "",
        "seconds": record.get("wall_seconds") or "",
        "tokens": (record.get("usage_total") or {}).get("total_tokens") or "",
        "needs_review": "是" if needs_review else "",
        "review_reason": review_reason,
    }


def classify_failure(status: str, errors: list[str], reason_parts: list[str]) -> str:
    """把失败归类成人能直接决策的几档——重点是区分"程序还能再试"和"得靠人"。

    背景：上一轮 424 行里 168 行统一标成"需人工确认"，但其中很大一部分其实是
    "搜索引擎没匹配到 → 换入口就能救"，把这类推给人工是浪费人最贵的资源。
    分类依据来自节点写进 errors 的文案（见 graph.search / rank）与状态。
    """
    if status.startswith("success"):
        return ""
    text = " ".join(errors + reason_parts)
    if status == "timeout" or "秒未结束" in text:
        return "超时（可再试）"
    if "头名词兜底页" in text or "与查询无关，已丢弃" in text:
        return "搜索引擎无匹配（需换入口）"
    if "第三方文库/内容站" in text:
        return "只搜到内容站（需换入口）"
    if "搜索引擎无匹配，改用官方入口" in text:
        return "已尝试官方入口仍未拿到"
    if "搜索无结果" in text:
        return "搜索无结果（索引缺失）"
    if status == "skipped_withdrawn":
        return "判为废止/失效（待人工确认）"
    if "下载失败" in text:
        return "下载环节失败"
    if status == "needs_confirmation":
        return "需人工确认"
    return ""


def review_flag(
    status: str,
    warnings: list[str],
    notes: list[str],
    errors: list[str],
    skip_reason: str,
    quarantined: bool,
    has_verification: bool,
) -> tuple[bool, str]:
    """人工复核标记：不只看成功与否，也看"成功了但可能有坑"的情况。"""
    reasons: list[str] = []
    if status in {"success_page_export"}:
        reasons.append("网页导出，非官方原件")
    if status in {"success_scanned"}:
        reasons.append("扫描件，无文本层，内容未能与名称比对")
    if status in {"not_downloadable", "needs_confirmation", "timeout", "error"}:
        reasons.append("未成功")
    if status == "skipped_withdrawn":
        reasons.append("判定为废止/失效，请人工确认是否仍要用")
    if quarantined:
        reasons.append("曾下到无关文件，已移入 _rejected")
    if any("文本层" in item for item in warnings + notes):
        reasons.append("扫描件（文本层不可读），名称未能核对")
    if any("内容未核对" in item for item in warnings + notes):
        # Word/表格等格式以前完全不做内容核对；现在抽不出文本时会留下这条提示，
        # 交付表里要能看见"这一件只核对了类型与来源"。
        reasons.append("内容未核对（该格式抽不出文本），建议抽查")
    if any("没有文件正文" in item for item in warnings):
        reasons.append("产物里没有文件正文（网页外壳/免费全文不可得）")
    if any("无关" in item or "不匹配" in item for item in warnings):
        reasons.append("内容与名称存在不一致提示")
    if not has_verification and status.startswith("success"):
        reasons.append("缺少校验明细")
    return bool(reasons), "；".join(reasons)


def build_workbook(source: Path, sheet_name: str, records: dict[str, dict[str, Any]], runs: list[str], output: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(source)
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"工作表不存在：{sheet_name}（现有：{workbook.sheetnames}）")
    sheet = workbook[sheet_name]

    header = [cell.value for cell in sheet[1]]
    base_columns = len(header)
    for offset, (title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=base_columns + offset, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(base_columns + offset)].width = width
    sheet.freeze_panes = sheet.cell(row=2, column=1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(base_columns + len(COLUMNS))}{sheet.max_row}"

    stats = Counter()
    failure_stats: Counter = Counter()
    review_rows: list[list[Any]] = []
    matched = 0
    for row_index in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row_index, column=1).value
        if name in (None, ""):
            continue
        record = records.get(normalize(str(name)))
        if not record:
            summary = {
                "label": "未爬取",
                "run_label": "",
                "status": "",
                "location": "",
                "filename": "",
                "check": "",
                "warnings": "",
                "reason": "",
                "notes": "",
                "standard_number": "",
                "sha256": "",
                "mime": "",
                "pages": "",
                "size": 0,
                "source": "",
                "seconds": "",
                "tokens": "",
                "failure_kind": "",
                "needs_review": "是",
                "review_reason": "本次批次未包含该行",
            }
        else:
            matched += 1
            summary = summarize(record, read_task_detail(record))
        stats[summary["label"]] += 1
        failure_stats[summary.get("failure_kind") or "（无）"] += 1

        # 按列名取值：列的顺序调整（例如插入「过程提示」）不会影响写入位置。
        row_values = {
            "行号": row_index,
            "爬取结果": summary["label"],
            "命中轮次": summary.get("run_label", ""),
            "失败原因": summary["reason"],
            "失败归类": summary.get("failure_kind", ""),
            "过程提示": summary["notes"],
            "文件位置": summary["location"],
            "文件名": summary["filename"],
            "校验结果": summary["check"],
            "校验提示": summary["warnings"],
            "识别到的文号": summary["standard_number"],
            "内容哈希(sha256)": summary["sha256"],
            "类型": summary["mime"],
            "页数": summary["pages"],
            "大小(KB)": round(summary["size"] / 1024, 1) if summary["size"] else "",
            "来源页面": summary["source"],
            "耗时(秒)": summary["seconds"],
            "token": summary["tokens"],
            "需人工复核": summary["needs_review"],
            "人工结论": "",
            "结果状态(原始)": summary["status"],
        }
        wrap_columns = {"失败原因", "校验结果", "校验提示", "过程提示", "来源页面"}
        status_column = 0
        for offset, (title, _width) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row_index, column=base_columns + offset, value=row_values.get(title, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=title in wrap_columns)
            if title == "爬取结果":
                status_column = base_columns + offset
        fill = STATUS_FILL.get(summary["status"], "FFF2CC" if summary["label"] == "未爬取" else None)
        if fill and status_column:
            sheet.cell(row=row_index, column=status_column).fill = PatternFill("solid", fgColor=fill)
        if summary["needs_review"] == "是":
            review_rows.append([row_index, str(name)[:60], summary["label"], summary["review_reason"], summary["location"], summary["source"]])

    build_summary_sheet(workbook, stats, failure_stats, runs, matched, sheet.max_row - 1)
    build_review_sheet(workbook, review_rows)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return {"stats": dict(stats), "matched": matched, "review": len(review_rows)}


def build_summary_sheet(workbook, stats: Counter, failure_stats: Counter, runs: list[str], matched: int, total: int) -> None:
    title = "汇总统计"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    sheet.append(["表格行数", total])
    sheet.append(["本次匹配到的行数", matched])
    sheet.append(["未爬取行数", total - matched])
    sheet.append([])
    sheet.append(["爬取结果", "行数"])
    for label, count in stats.most_common():
        sheet.append([label, count])
    sheet.append([])
    # 失败归类：把"程序还能再试"和"得靠人"分开，人工复核时按这个列排序即可。
    sheet.append(["失败归类", "行数"])
    for kind, count in failure_stats.most_common():
        sheet.append([kind, count])
    sheet.append([])
    sheet.append(["数据来源批次", ""])
    for item in runs:
        sheet.append(["", item])
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 60
    for row in sheet.iter_rows(min_row=6, max_row=6):
        for cell in row:
            cell.font = Font(bold=True)


def build_review_sheet(workbook, rows: list[list[Any]]) -> None:
    title = "问题清单"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(["行号", "规范/文件名称", "爬取结果", "需复核原因", "文件位置", "来源页面"])
    for row in rows:
        sheet.append(row)
    widths = [6, 52, 16, 40, 34, 44]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(prog="official-file-report", description="把爬取结果汇总回编制依据表，输出带新增列的 Excel")
    parser.add_argument("--source", default=str(DEFAULT_SOURCE), help="原始 Excel 路径")
    parser.add_argument("--sheet", default=SHEET_NAME, help="工作表名")
    parser.add_argument("--results", nargs="*", help="批次产物目录（可多个），默认汇总 results/batch_sheet_*")
    parser.add_argument("--output", help="输出 Excel 路径，默认与原表同目录、文件名加后缀 _爬取结果")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    source = Path(args.source)
    if not source.exists():
        raise SystemExit(f"原始表格不存在：{source}")

    if args.results:
        roots = [Path(item) for item in args.results]
    else:
        roots = sorted(Path(".").glob(DEFAULT_RESULT_GLOB))
    if not roots:
        raise SystemExit("没有找到任何批次产物目录（results/batch_sheet_*）")

    records, runs = load_run(roots)
    output = Path(args.output) if args.output else source.with_name(f"{source.stem}_爬取结果_{datetime.now():%Y%m%d_%H%M}.xlsx")
    result = build_workbook(source, args.sheet, records, runs, output)

    print(f"批次目录：{len(roots)} 个")
    for item in runs:
        print(f"  - {item}")
    print(f"可匹配任务：{len(records)} 条")
    print(f"写入表格行：{result['matched']} 行匹配成功")
    print("结果分布：")
    for label, count in sorted(result["stats"].items(), key=lambda pair: -pair[1]):
        print(f"  {label}: {count}")
    print(f"需人工复核：{result['review']} 行（见「问题清单」工作表）")
    print(f"输出：{output}")


if __name__ == "__main__":
    main()
