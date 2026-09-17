"""从编制依据汇总表读取任务，按行驱动批量爬取。

用法（在项目根目录、带 PYTHONPATH=src 运行）：
    python -m standards_crawler_agent.sheet_crawl --list-types           # 先看分类与抽样
    python -m standards_crawler_agent.sheet_crawl --rows 3,4,5 --jobs 3  # 指定行号试点
    python -m standards_crawler_agent.sheet_crawl --remark 下载 --limit 10
    python -m standards_crawler_agent.sheet_crawl --start 1 --end 100 --jobs 3 --timeout 300

产物：results/batch_sheet_<时间戳>/<序号_名称>/{result.json,run.log} + batch_summary.json
（序号是**去重后**的位置，所以汇总回表格时按名称匹配，见 report.py）
"""

from __future__ import annotations

import argparse
import os
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

# 编制依据表是各人自己的业务文件，不进版本库，所以默认值只写**相对路径**（相对当前工作目录）。
# 用环境变量 SHEET_SOURCE / SHEET_NAME 覆盖（推荐写进 .env），或用命令行 --source / --sheet 覆盖。
DEFAULT_SOURCE = Path(os.getenv("SHEET_SOURCE", "编制依据汇总表.xlsx"))
SHEET_NAME = os.getenv("SHEET_NAME", "编制依据")


@dataclass
class Row:
    index: int  # 表内行号（含表头，从 1 数起，与 Excel 显示一致）
    name: str
    number: str
    remark: str


def read_rows(source: Path, sheet_name: str) -> list[Row]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"工作表不存在：{sheet_name}（现有：{workbook.sheetnames}）")
    sheet = workbook[sheet_name]
    rows: list[Row] = []
    for position, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        if position == 1:
            continue
        name = str(values[0]).strip() if values and values[0] is not None else ""
        if not name:
            continue
        number = str(values[1]).strip() if len(values) > 1 and values[1] is not None else ""
        remark = str(values[2]).strip() if len(values) > 2 and values[2] is not None else ""
        rows.append(Row(index=position, name=name, number=number, remark=remark))
    return rows


def classify(number: str) -> str:
    """按文号粗略分类，用于抽样时保证文种覆盖。"""
    value = number.replace(" ", "")
    rules = [
        ("GB", "国家标准"),
        ("JGJ", "建工行业标准"),
        ("JTG", "交通行业标准"),
        ("CJJ", "城建行业标准"),
        ("DB", "地方标准"),
        ("YB", "冶金行业标准"),
        ("主席令", "法律"),
        ("国务院令", "行政法规"),
        ("住房和城乡建设部令", "部门规章"),
        ("人民政府令", "地方政府规章"),
        ("济建", "济南市文件"),
        ("鲁建", "山东省文件"),
        ("建质", "住建部规范性文件"),
        ("建办质", "住建部规范性文件"),
    ]
    for prefix, label in rules:
        if prefix in value:
            return label
    return "其他"


def select(rows: list[Row], args: argparse.Namespace) -> list[Row]:
    chosen = rows
    if args.rows:
        wanted = {int(item) for item in args.rows.split(",") if item.strip()}
        chosen = [row for row in chosen if row.index in wanted]
    if args.remark:
        chosen = [row for row in chosen if row.remark == args.remark]
    if args.start:
        chosen = [row for row in chosen if row.index >= args.start]
    if args.end:
        chosen = [row for row in chosen if row.index <= args.end]
    if args.spread:
        # 按文种轮转抽取，保证试点里各种文号类型都有一条。
        # 不传 --limit 时取满全部选中行（此前漏了这个兜底，会拿 None 去比大小直接报错）。
        limit = args.limit or len(chosen)
        buckets: dict[str, list[Row]] = {}
        for row in chosen:
            buckets.setdefault(classify(row.number), []).append(row)
        picked: list[Row] = []
        while len(picked) < limit and any(buckets.values()):
            for key in sorted(buckets):
                if buckets[key] and len(picked) < limit:
                    picked.append(buckets[key].pop(0))
        return sorted(picked, key=lambda row: row.index)
    if args.limit:
        chosen = chosen[: args.limit]
    return chosen


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(prog="official-file-sheet-crawl", description="按编制依据表逐行爬取")
    parser.add_argument("--source", default=str(DEFAULT_SOURCE))
    parser.add_argument("--sheet", default=SHEET_NAME)
    parser.add_argument("--rows", help="指定表内行号，逗号分隔，例如 3,4,5")
    parser.add_argument("--remark", help="只跑备注等于该值的行，例如 下载")
    parser.add_argument("--start", type=int, help="起始行号（含）")
    parser.add_argument("--end", type=int, help="结束行号（含）")
    parser.add_argument("--limit", type=int, help="最多跑多少条")
    parser.add_argument("--spread", action="store_true", help="按文种轮转抽样（与 --limit 搭配）")
    parser.add_argument("--list-types", action="store_true", help="只打印分类统计与抽样预览，不爬取")
    parser.add_argument("--jobs", type=int, default=3)
    parser.add_argument("--timeout", type=int, default=300)
    parser.add_argument("--retries", type=int, default=1)
    parser.add_argument("--stagger", type=float, default=2.0)
    parser.add_argument("--output-root", help="批次产物目录，默认 results/batch_sheet_<时间戳>")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    rows = read_rows(Path(args.source), args.sheet)

    if args.list_types:
        counts: dict[str, int] = {}
        samples: dict[str, list[Row]] = {}
        for row in rows:
            label = classify(row.number)
            counts[label] = counts.get(label, 0) + 1
            samples.setdefault(label, []).append(row)
        print(f"共 {len(rows)} 行")
        for label, count in sorted(counts.items(), key=lambda pair: -pair[1]):
            print(f"\n== {label}：{count} 条")
            for row in samples[label][:3]:
                print(f"   {row.index:>4} | {row.name[:52]:<54} | {row.number} | 备注={row.remark or '-'}")
        return

    chosen = select(rows, args)
    if not chosen:
        raise SystemExit("没有选中任何行")
    print(f"选中 {len(chosen)} 行（表内行号：{chosen[0].index}~{chosen[-1].index}）")
    for row in chosen:
        print(f"  {row.index:>4} | {row.name[:56]:<58} | {row.number} | 备注={row.remark or '-'}")

    output_root = args.output_root or f"results/batch_sheet_{__import__('datetime').datetime.now():%Y%m%d_%H%M%S}"
    from .batch import main as batch_main

    batch_main(
        [row.name for row in chosen]
        + [
            "--jobs", str(args.jobs),
            "--stagger", str(args.stagger),
            "--timeout", str(args.timeout),
            "--retries", str(args.retries),
            "--output-root", output_root,
        ]
    )


if __name__ == "__main__":
    main()
