from __future__ import annotations

import hashlib
import mimetypes
import re
import zipfile
from pathlib import Path

import httpx
from pypdf import PdfReader

from .content import (
    detect_absent_body,
    detect_scanned_only,
    extract_standard_number,
    normalize_standard_number,
)
from .models import Verification


WITHDRAWN_TERMS = ("已废止", "已作废", "废止", "作废", "失效", "停止执行", "不再执行")
REPLACEMENT_TERMS = ("替代", "代替", "修订后", "新标准", "新版本", "同时发布")
ACTIVE_TERMS = ("未废止", "未作废", "仍然有效", "继续有效", "现行有效")
# 指向自身的有效信号。发布类词（“现予以发布”“自…起施行”）同样属于自我有效：
# 批准发布公告里必然带着“原…同时废止”，但被废止的是旧版本。
SELF_ACTIVE_TERMS = (
    "现行",
    "现行有效",
    "未废止",
    "未作废",
    "仍然有效",
    "继续有效",
    "予以发布",
    "批准发布",
    "批准为",
    "起施行",
)
SELF_WITHDRAWN_TERMS = ("已废止", "已作废", "本文件废止", "本标准废止", "停止执行", "不再执行")


def _strip_brackets(value: str) -> str:
    return re.sub(r"[《》〈〉「」『』]", "", value or "").strip()


def _subject_windows(text: str, subject: str, window: int) -> str:
    """只保留目标名称出现位置附近的文本。

    法律法规汇编页 / 列表页会在同一页列出多份文件的存废状态，
    直接扫描整页会把“别的文件停止执行”误判成本文件已废止
    （实测：应急管理部页面上《生产安全事故应急条例》被误判为停止执行）。
    命中不到目标名称时退回整页文本，保持原有行为。
    """
    if not subject:
        return text
    spans: list[str] = []
    start = 0
    while len(spans) < 8:
        index = text.find(subject, start)
        if index < 0:
            break
        spans.append(text[max(0, index - window) : index + len(subject) + window])
        start = index + len(subject)
    return " ".join(spans) if spans else text


def _snippet_withdrawal(snippet: str, subject: str, proximity: int = 60) -> str | None:
    """搜索结果摘要里“名称 + 存废词”紧邻出现时，视为明确的存废标记。"""
    if not snippet or not subject:
        return None
    index = snippet.find(subject)
    if index < 0:
        return None
    scope = snippet[max(0, index - proximity) : index + len(subject) + proximity]
    for term in WITHDRAWN_TERMS + SELF_WITHDRAWN_TERMS:
        if term in scope:
            return term
    return None


def detect_lifecycle_status(page_text: str = "", *, subject: str = "", snippet: str = "", window: int = 120) -> tuple[str, str | None]:
    """判断目标文件是否已废止。

    判定依据（顺序即优先级）：
      1. 页面正文里**目标名称附近**的文本——避免汇编页/首页上其它文件的存废状态误伤
         （实测：应急管理部页面正文第二十五条里的“停止执行”被误当成本法规废止）；
      2. 搜索结果摘要里“名称 + 存废词”紧邻出现——列表页的摘要本身就带状态标记；
      3. 正文里压根没出现目标名称时**不做废止判定**——那个页面大概率不是目标页面
         （实测：落到政府门户首页时，首页上别的文件的“废止”会误伤目标名称），
         此时应由决策环节判断“这页没用”，而不是断言目标文件已废止。
    """
    subject = _strip_brackets(subject)
    snippet = _strip_brackets(snippet or "")
    combined = (page_text or "").strip()

    if subject and subject not in combined:
        term = _snippet_withdrawal(snippet, subject)
        if term:
            return "withdrawn", f"搜索结果摘要显示该文件{term}，跳过下载"
        return "active", "（页面正文未出现目标名称，未做废止判定）"

    scoped = _subject_windows(combined, subject, window)
    scope_note = ""
    if scoped != combined:
        scope_note = "（判定范围限定在目标名称附近）"
        # 整页有废止信号但都不在目标名称附近时，把这件事显式写出来，
        # 便于人工判断是不是又一次“汇编页误伤”。
        page_signals = [term for term in WITHDRAWN_TERMS if term in combined]
        page_signals += [
            term for term in SELF_WITHDRAWN_TERMS if term in combined and term not in page_signals
        ]
        if page_signals and not any(term in scoped for term in page_signals):
            scope_note = f"（整页出现废止信号“{page_signals[0]}”但不在目标名称附近，已忽略）"
    active = [term for term in ACTIVE_TERMS if term in scoped]
    withdrawn = [term for term in WITHDRAWN_TERMS if term in scoped]
    replacements = [term for term in REPLACEMENT_TERMS if term in scoped]
    # 实验性修复：标准详情页会把“标准状态：现行”与“代替了以下标准 / 被代替”混排，
    # 单独出现“废止”字样不等于当前文件已废止，因此先看指向自身的状态表述。
    self_active = [term for term in SELF_ACTIVE_TERMS if term in scoped]
    self_withdrawn = [term for term in SELF_WITHDRAWN_TERMS if term in scoped]
    if self_active and not self_withdrawn:
        return "active", f"页面自身状态明确为有效（{', '.join(self_active)}）{scope_note}"
    if active and not any(term in scoped.replace("未废止", "").replace("未作废", "") for term in ("废止", "作废")):
        return "active", f"页面明确说明文件仍有效（{', '.join(active)}）"
    if withdrawn and replacements:
        return "ambiguous", f"页面同时出现废止信号（{', '.join(withdrawn)}）和新版本/替代说明，未自动下载{scope_note}"
    if withdrawn:
        return "withdrawn", f"页面明确标记文件{withdrawn[0]}，跳过下载{scope_note}"
    # 窗口内没有任何存废信号时也不丢信息：把“整页信号已被忽略”的说明带出去。
    return "active", scope_note or None


def safe_filename_part(value: str) -> str:
    """清理 Windows 不允许出现在文件名中的字符，同时保留可读性。"""
    value = value.replace("/", "／").replace("\\", "＼")
    value = re.sub(r'[<>:"|?*]', "_", value)
    return re.sub(r"\s+", " ", value).strip(" .")


# extract_standard_number / normalize_standard_number 已迁到 content.py（单一owner），
# 见本文件顶部的 import。


def acceptable_title(candidate: str | None, requested_name: str) -> str | None:
    """对模型给出的正式名称做体检。

    模型有时会把手误填成站点名或栏目标题（“国家标准 - 全国标准信息公共服务平台”），
    这类名称不能拿来当文件名；与请求名重合度过低的也一律不采用。
    """
    if not candidate:
        return None
    value = re.sub(r"\s+", " ", candidate).strip(" ._-—–")
    if not 4 <= len(value) <= 80:
        return None
    if any(word in value for word in ("首页", "登录", "网站", "平台", "门户", "百科", "导航")):
        return None
    requested_chars = {char for char in re.sub(r"[\s《》()（）\-_/]", "", requested_name)}
    if requested_chars:
        overlap = len(set(value) & requested_chars) / len(requested_chars)
        if overlap < 0.5:
            return None
    return value


def build_canonical_filename(
    requested_name: str,
    extension: str,
    standard_number: str | None,
    matched_title: str | None = None,
) -> str:
    """生成 名称_编号.扩展名。

    名称优先用页面上识别到的正式名称：编制依据表里的名称常有笔误
    （把“混凝土结构工程施工质量验收规范”写成“混凝土结构施工质量验收规范”），
    用正式名称落盘才能和标准原文对得上；模型没给出或体检不通过时回退到请求名。
    """
    base = acceptable_title(matched_title, requested_name) or requested_name
    base = re.sub(r"\.(pdf|docx?|xlsx?|zip)$", "", base.strip(), flags=re.I)
    base = safe_filename_part(base)
    number = normalize_standard_number(standard_number)
    if number:
        return f"{base}_{safe_filename_part(number)}{extension.lower()}"
    return f"{base}{extension.lower()}"


def download_direct(url: str, output_dir: str, max_bytes: int = 100 * 1024 * 1024) -> str:
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    name = url.rsplit("/", 1)[-1].split("?", 1)[0] or "downloaded_file"
    path = Path(output_dir) / name
    try:
        with httpx.stream("GET", url, follow_redirects=True, timeout=60) as response:
            response.raise_for_status()
            content_type = (response.headers.get("content-type") or "").lower()
            if "text/html" in content_type:
                raise RuntimeError(
                    f"目标地址返回的是网页而不是文件（content-type: {content_type}）——直链不能指向网页，请改用页面导出或附件链接"
                )
            total = 0
            with path.open("wb") as handle:
                for chunk in response.iter_bytes():
                    total += len(chunk)
                    if total > max_bytes:
                        raise ValueError("file exceeds configured size limit")
                    handle.write(chunk)
    except Exception:
        # 中途失败不要留下半截文件（它会和被拒产物一样混进下载目录）。
        remove_partial(str(path))
        raise
    return str(path)


# 部分政府站点的附件 URL 不带扩展名（例如详情页里指向 .../art_17339_761192.html 的附件），
# 落盘后文件名没有后缀。仅靠文件名既判断不出类型、也无法进入 PDF 解析分支，
# 因此这里按文件头识别类型，并在重命名时补回扩展名。
ZIP_CONTAINERS: tuple[tuple[str, str, str], ...] = (
    ("word/", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    ("xl/", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"),
    ("ppt/", "application/vnd.openxmlformats-officedocument.presentationml.presentation", ".pptx"),
    ("OFD.xml", "application/ofd", ".ofd"),
)


def sniff_file_type(path: Path) -> tuple[str, str]:
    """按文件头判断类型，返回 (mime, 扩展名)；无法识别时两者均为空字符串。"""
    try:
        with path.open("rb") as handle:
            head = handle.read(8)
    except OSError:
        return "", ""
    if head.startswith(b"%PDF-"):
        return "application/pdf", ".pdf"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
        except Exception:
            return "application/zip", ".zip"
        for marker, mime, suffix in ZIP_CONTAINERS:
            if any(name == marker or name.startswith(marker) for name in names):
                return mime, suffix
        return "application/zip", ".zip"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        # 老式 OLE2 复合文档，doc/xls/ppt 共用同一文件头，无法仅凭文件头区分扩展名。
        return "application/x-ole-storage", ""
    if head.startswith(b"{\\rtf"):
        return "application/rtf", ".rtf"
    return "", ""


# 已知的文档扩展名。扩展名不在这个集合里（例如政府站下载端点直接落盘的 “.jsp”）时，
# 一律以文件头判定的类型为准：实测山东省政府站的 .../module/download/downfile.jsp
# 返回的是真正的 PDF，但文件名叫 *.jsp，旧逻辑只按扩展名判断类型，于是跳过了 PDF
# 内容校验，把一份只有 4 页的打印片段当成成功结果（title_matches 为空、零 warning）。
DOCUMENT_SUFFIXES = frozenset(
    {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".ofd", ".rtf", ".zip", ".rar"}
)


def resolve_file_type(path: Path) -> tuple[str, str, str]:
    """判定文件的真实类型，返回 (mime, 应当采用的扩展名, 异常说明)。

    扩展名缺失、不在已知文档扩展名里、或与文件头判定冲突时，以文件头为准，
    并把应当采用的扩展名一并返回，供校验分支与落盘重命名使用。
    """
    suffix = path.suffix.lower()
    declared_mime = mimetypes.guess_type(path.name)[0] or ""
    sniffed_mime, sniffed_suffix = sniff_file_type(path)
    if not sniffed_mime:
        return declared_mime, "", ""
    unknown_suffix = bool(suffix) and suffix not in DOCUMENT_SUFFIXES
    conflict = bool(suffix) and bool(sniffed_suffix) and suffix != sniffed_suffix
    if not suffix or unknown_suffix or conflict:
        if not suffix:
            return sniffed_mime, sniffed_suffix, ""
        actual = sniffed_suffix or sniffed_mime
        note = f"文件扩展名 {suffix} 与实际类型不符（实际为 {actual}），已按实际类型校验与命名"
        return sniffed_mime, sniffed_suffix, note
    return declared_mime or sniffed_mime, "", ""


# 网页打印稿的页码标记（浏览器打印页脚形如 “14/17”）。
PRINT_PAGE_MARKER = re.compile(r"(?<![\d/])(\d{1,3})\s*/\s*(\d{1,4})(?![\d/])")


def detect_print_fragment(text: str, page_count: int, min_total: int = 5) -> str | None:
    """检测“网页打印片段”：文件自带页码标记说明它只是一份更长文档的一部分。

    实测：山东省政府站下载端点给的 PDF 只有 4 页，页脚却反复出现 14/17、15/17、16/17，
    正文（第一章…第七章）整段缺失——而页眉恰好带着完整标题，所以按名称匹配反而能通过。
    判据是同一个分母重复出现（打印稿每页都有页脚）且大于文件实际页数。
    """
    if page_count <= 0:
        return None
    totals: dict[int, int] = {}
    samples: dict[int, tuple[int, int]] = {}
    for match in PRINT_PAGE_MARKER.finditer(text or ""):
        current, total = int(match.group(1)), int(match.group(2))
        if not (1 <= current < total <= 300) or total < min_total:
            continue
        totals[total] = totals.get(total, 0) + 1
        samples.setdefault(total, (current, total))
    for total, count in totals.items():
        if count >= 2 and total > page_count:
            current, _ = samples[total]
            return (
                f"疑似网页打印片段：文中页码标记显示共 {total} 页（如 {current}/{total}），"
                f"但文件只有 {page_count} 页，正文可能不完整"
            )
    return None


def check_claimed_title(claimed: str | None, requested_name: str) -> tuple[bool, str]:
    """校验模型自己声称的“目标文件名称”是否与请求名称有关。

    只对“明显无关”判失败：把手误填成站点名/栏目标题（“国家标准 - 全国标准信息公共服务平台”）
    属于常见噪声，不足以判定下错文件；但字符重合度低于 50% 时，说明模型认的目标
    根本不是这份文件——实测依据：79 次填了 matched_title 的决策里只有 6 次重合度不达标，
    6 次都是同一份下错的文件（山东省政府站公报汇总页上另一份公文的附件），却都报了 success。
    """
    if not claimed:
        return True, ""
    value = re.sub(r"\s+", " ", claimed).strip(" ._-—–")
    requested_chars = {char for char in re.sub(r"[\s《》()（）\-_/]", "", requested_name)}
    if not requested_chars:
        return True, ""
    overlap = len(set(value) & requested_chars) / len(requested_chars)
    if overlap < 0.5:
        return False, f"模型声称的目标文件名称与请求名称无关（“{value}”，重合度 {overlap:.0%}），判定为下错文件"
    return True, ""


def sha256_of(path: Path, chunk_size: int = 1024 * 1024) -> str:
    """分块计算 sha256，避免把几十 MB 的文件一次性读进内存。"""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def same_content(left: Path, right: Path) -> bool:
    """两个文件内容是否相同（先比大小，再比 sha256）。

    用途：产物落盘后要改成规范名（`名称_编号.扩展名`），而目标名**可能已经存在**。
    同名文件的处理规则是"新版直接盖旧版"，但**内容一模一样时没必要重写**：
    重写一次既白花一次几十 MB 的写盘，又会把旧文件的时间戳改掉（看不出它其实没变）。
    所以这里只做一件事——判断"是不是同一个字节序列"。

    哈希会读整个文件，因此**先用大小筛**：绝大多数不同内容的文件大小就不同。
    """
    try:
        if left.stat().st_size != right.stat().st_size:
            return False
    except OSError:
        return False
    return sha256_of(left) == sha256_of(right)


# 被校验拒绝的产物统一移到这里，不留在下载目录里冒充结果。
REJECTED_DIRNAME = "_rejected"


def quarantine_artifact(path: str | None, reason: str, download_dir: str = "downloads") -> str | None:
    """把被校验拒绝的产物移出下载目录，返回新路径；无法处理时返回 None。

    实测：被拒的文件会一直留在 downloads/ 里混进结果（r15/r16 的 89d65074….pdf、
    r17 的 downfile.jsp）。移到 <download_dir>/_rejected/ 后既不污染结果目录，
    又能事后排查——原因写在同名 .reason.txt 里。
    """
    if not path:
        return None
    source = Path(path)
    if not source.is_file():
        return None
    target_dir = Path(download_dir) / REJECTED_DIRNAME
    try:
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / source.name
        if target.exists():
            # 多轮回归里同一个下错的文件名会反复出现，用内容摘要区分，不覆盖旧记录。
            target = target_dir / f"{source.stem}__{sha256_of(source)[:8]}{source.suffix}"
        source.replace(target)
        target.with_suffix(target.suffix + ".reason.txt").write_text(reason, encoding="utf-8")
    except OSError:
        return None
    return str(target)


def remove_partial(path: str | None) -> None:
    """删除下到一半的文件（网络中断、超过体积上限等），避免半截文件留在下载目录。"""
    if not path:
        return
    try:
        Path(path).unlink(missing_ok=True)
    except OSError:
        pass


# 高频汉字：乱码文本层里几乎不会出现它们，用来判断文本层能不能读。
COMMON_HANZI = set(
    "的一是在不了有和人这中大为上个国我以要他时来用们工程管理规定通知要求加强标准实施"
    "建设局省市区县房屋建筑安全危险较大分部项管理发布年月日第条文件编号"
)


def looks_readable_text(text: str, min_ratio: float = 0.01, min_length: int = 120) -> bool:
    """判断 PDF 文本层是否可读。

    扫描件、字体编码异常的文件会抽出乱码（实测《工程结构通用规范》首页抽出
    形如 “OOb?W蜰a^鷭緪鐿醏oQl_” 的内容），这种文本层不能用来判断内容是否匹配。
    """
    stripped = [char for char in text if not char.isspace()]
    if len(stripped) < min_length:
        return False
    common = sum(1 for char in stripped if char in COMMON_HANZI)
    return common / len(stripped) >= min_ratio


def name_segments(name: str, min_length: int = 4) -> list[str]:
    """把请求名称切成有辨识度的片段，用于校验下载内容。

    实测反例：请求“济南市住房和城乡建设局关于进一步加强全市房屋建筑和燃气供热工程
    危险性较大分部分项工程管理的通知”，下到的却是该局的另一份文件
    （济建建管字〔2020〕48号，关于项目交易文件在线监督管理）。
    两份文件的差异恰好在“危险性较大分部分项工程”这类片段上，
    所以要按片段校验，而不是按单字覆盖率。
    """
    value = re.sub(r"\.(pdf|docx?|xlsx?|zip)$", "", name.strip(), flags=re.I)
    value = re.sub(r"[《》（）()、，,。；;：:·\s]", "", value)
    for stop in ("关于印发", "关于", "进一步加强", "进一步", "全市", "的通知", "通知", "的", "印发"):
        value = value.replace(stop, "|")
    return [part for part in value.split("|") if len(part) >= min_length]


def name_char_coverage(text: str, requested_name: str) -> float:
    """请求名称中的字符在文本里的覆盖比例。

    用来容忍请求名本身的笔误：实测请求“混凝土结构施工质量验收规范”（漏了“工程”），
    而官方正式名是“混凝土结构工程施工质量验收规范”，按整段匹配会 0 命中，
    按字符覆盖则是 13/13。
    """
    wanted = {
        char
        for char in re.sub(r"[\s《》（）()、，,。；;：:·\-_/]", "", requested_name)
        if char.strip()
    }
    if not wanted:
        return 0.0
    return len(wanted & set(text)) / len(wanted)


def check_name_match(
    text: str,
    requested_name: str,
    min_ratio: float = 0.6,
    min_coverage: float = 0.9,
) -> tuple[bool, str]:
    """校验下载内容是否确实是目标文件。

    两条判据满足其一即通过：
      - 名称片段命中率 ≥ 60%（名称准确时最可靠）
      - 名称字符覆盖率 ≥ 90%（名称有笔误、漏字时兜底）

    再加一条**否决**：产物写的是"同名的另一份地方文件"时，无论上面两条怎么算都要拒
    （见 detect_near_miss_document）。

    只在文本层可读时判断；乱码或过短的文本层直接放行，避免误伤正确文件。
    实测的误杀案例：把正确的 140 页标准判成不匹配后换候选，下到了 75 页扫描件——比原来更差。
    """
    if not looks_readable_text(text):
        return True, "文本层不可读或过短，跳过名称校验"
    segments = name_segments(requested_name)
    if not segments:
        return True, "名称无法切分出有效片段，跳过校验"
    near_miss = detect_near_miss_document(text, requested_name)
    if near_miss:
        return False, near_miss
    # **比片段时忽略空白**：Word/PDF 抽出来的标题常被换行切开（实测《危险性较大的分部分项工程
    # 专项施工方案编制指南》的标题在 docx 里是"专项施工方案\n编制指南"），不掐掉空白就会把
    # 正确文件判成"片段命中 1/2"。注意：这条只在**校验**场景成立——排序场景里掐空白会反过来
    # 抬高转载页（见 ranking.page_score 的注释），两处目标不同，结论也不同。
    squashed = re.sub(r"[\s\u3000]+", "", text)
    hits = [segment for segment in segments if re.sub(r"[\s\u3000]+", "", segment) in squashed]
    ratio = len(hits) / len(segments)
    coverage = name_char_coverage(text, requested_name)
    detail = f"名称片段命中 {len(hits)}/{len(segments)}，字符覆盖 {coverage:.0%}"
    if ratio >= min_ratio or coverage >= min_coverage:
        return True, detail
    return False, f"下载内容与目标名称不匹配（{detail}）"


# 文种词分族：只有**技术标准类**允许互换（实测把《高层民用建筑钢结构技术规范》写成
# "…技术规程"必须放行，"规范/规程/标准"是同一份文件的不同叫法）；规章类每个词各自成族——
# "规定"和"办法"背后往往就是两份不同的文件。
TYPE_WORD_FAMILY = {
    "技术规范": "标准", "技术规程": "标准", "规范": "标准", "规程": "标准",
    "标准": "标准", "通则": "标准", "规则": "标准", "技术要求": "标准",
    "实施细则": "细则", "细则": "细则",
    "实施办法": "办法", "办法": "办法",
    "规定": "规定", "条例": "条例", "通知": "通知", "公告": "通知", "通报": "通知",
    "意见": "意见", "方案": "方案", "导则": "导则", "指南": "指南", "决定": "决定",
}
PROVINCE_NAMES = re.compile(
    r"(?:北京|上海|天津|重庆|河北|山西|辽宁|吉林|黑龙江|江苏|浙江|安徽|福建|江西|山东|河南|湖北|湖南|"
    r"广东|海南|四川|贵州|云南|陕西|甘肃|青海|内蒙古|广西|西藏|宁夏|新疆|香港|澳门|台湾)"
)
_PUNCTUATION = re.compile(r"[\s\u3000《》〈〉「」『』（）()、，,。；;：:·\-—_/／－]")


def _type_word(name: str) -> str | None:
    """取名称里的文种词；命中多个时取最长的（"实施细则"优先于"细则"）。"""
    matches = [word for word in TYPE_WORD_FAMILY if word in (name or "")]
    return max(matches, key=len) if matches else None


def _core_name(name: str) -> str:
    """去掉标点与**文种词**之后的"核心名称"。"""
    value = _PUNCTUATION.sub("", name or "")
    word = _type_word(value)
    if word:
        index = value.rfind(word)
        value = value[:index] + value[index + len(word):]
    return value


def detect_near_miss_document(text: str, requested_name: str) -> str | None:
    """产物写的是"同名的另一份地方文件"时给出说明，否则返回 None。

    实测（r21 #3）：请求《建设工程施工现场管理规定》（建设部令第15号），产物却是
    《北京市建设工程施工现场管理办法》（北京市人民政府第247号令）。两者**核心名完全相同**，
    只差"北京市"和"规定→办法"两处，于是"名称片段命中 0/1"被"字符覆盖 100%"的兜底放行，
    交出去一份错文件。

    判据刻意做窄——三条必须同时成立：
      1. 请求名的**核心名**在产物文本里连续出现；
      2. 紧跟其后的必须是**另一个文种词**（不同族，例如"规定"的位置写着"办法"）；
      3. 核心名前面带**行政区划**（北京市 / 济南市 / 朝阳区…）。
    因此正确件不受影响：湖南住建厅转载的《建设工程施工现场管理规定》后面仍是"规定"；
    漏字笔误也不受影响：《混凝土结构施工质量验收规范》的核心名根本不会连续出现在
    正式名《混凝土结构工程施工质量验收规范》里，第 1 条就不成立。
    """
    squashed = re.sub(r"[\s\u3000]+", "", text or "")
    request_type = _type_word(requested_name)
    core = _core_name(requested_name)
    if not request_type or len(core) < 4:
        return None
    position = squashed.find(core)
    if position < 0:
        return None
    tail = squashed[position + len(core): position + len(core) + 6]
    artifact_type = _type_word(tail[:4])
    if not artifact_type or not tail.startswith(artifact_type):
        return None
    if TYPE_WORD_FAMILY.get(artifact_type) == TYPE_WORD_FAMILY.get(request_type):
        return None
    head = squashed[max(0, position - 6):position]
    province = PROVINCE_NAMES.search(head)
    if province:
        start = position - len(head) + province.start()
    elif re.search(r"[市县区州盟旗]$", head):
        # 只往回带两个汉字，做成"济南市""朝阳区"这样的窗口
        start = max(0, position - 3)
    else:
        return None
    window = squashed[start: position + len(core) + len(artifact_type)]
    return (
        f"产物是另一份文件（地方版本）：正文里写的是「{window}」，"
        f"而请求名的文种词是「{request_type}」——两者核心名相同、只差行政区划与文种词"
    )


# --------------------------------------------------------------------------
# 正文实质校验（判据本身在 content.py，这里只负责把 PDF 拆成文本与图像信息）
# --------------------------------------------------------------------------

# 名称校验沿用原来的 5 页窗口（改动它会影响既有判定），
# 正文实质校验则看更宽的窗口：目次、前言、条款常常不在前 5 页。
NAME_SAMPLE_PAGES = 5
BODY_SAMPLE_PAGES = 20


def count_image_pages(reader, sample: int) -> int:
    """统计前 sample 页里有多少页带图像 XObject（扫描件的判据之一）。"""
    count = 0
    for index in range(min(len(reader.pages), sample)):
        try:
            resources = reader.pages[index].get("/Resources")
            if resources is None:
                continue
            xobjects = resources.get_object().get("/XObject")
            if xobjects is None:
                continue
            # 把 XObject 字典解析一次就够：某些扫描件单页有几百个图像对象，
            # 在循环里反复 get_object() 会让整个校验慢到不可用。
            mapping = xobjects.get_object()
            for key in list(mapping.keys()):
                try:
                    if mapping[key].get_object().get("/Subtype") == "/Image":
                        count += 1
                        break
                except Exception:
                    continue
        except Exception:
            continue
    return count


def extract_page_text(reader, pages: int) -> str:
    """按页抽取文本，单页失败不影响整份文档。"""
    chunks: list[str] = []
    for index in range(min(len(reader.pages), pages)):
        try:
            chunks.append(reader.pages[index].extract_text() or "")
        except Exception:
            chunks.append("")
    return "\n".join(chunks)


# 非 PDF 产物（Word 文档）没有"页数"，用这个值让 detect_absent_body 跳过"篇幅极小"那一条
# ——那条判据依赖真实页数，套在 Word 上会误伤短公文。
WORD_PAGE_HINT = 99


def extract_docx_text(path: Path) -> str:
    """抽取 .docx 正文（段落 + 表格）。python-docx 是已安装的依赖。"""
    import docx  # 延迟导入：只有真的遇到 .docx 才需要它

    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(part for part in parts if part)


def extract_doc_text(path: Path) -> str:
    """抽取老式 .doc（OLE2）里的 WordDocument 流文本。

    Word 97-2003 的正文在 WordDocument 流中，可能是 UTF-16LE，也可能是按本地代码页
    压成 8 位的形态；两种都试一遍再拼起来——反正后续判据都是"子串是否存在"，
    命中任一种编码即可。**不做分片表解析**：只要够判定"有没有目标名称"就行。
    """
    import olefile  # 延迟导入

    if not olefile.isOleFile(str(path)):
        return ""
    with olefile.OleFileIO(str(path)) as ole:
        if not ole.exists("WordDocument"):
            return ""
        raw = ole.openstream("WordDocument").read()
    text = raw.decode("utf-16-le", errors="ignore")
    try:
        text += "\n" + raw.decode("gbk", errors="ignore")
    except Exception:
        pass
    return text


def extract_document_text(path: Path, kind: str) -> str:
    """按扩展名抽取 Word 文档正文；抽不出时返回空字符串（由调用方如实记录）。"""
    if kind == ".docx":
        return extract_docx_text(path)
    if kind == ".doc":
        return extract_doc_text(path)
    return ""


def verify_file(path: str, requested_name: str, claimed_title: str | None = None) -> Verification:
    """校验下载结果。

    claimed_title 是模型自己填的“页面上识别到的正式文件名称”。它与请求名完全无关时
    （实测：山东省政府站公报汇总页上另一份公文的附件），说明模型认错了目标文件。
    """
    file_path = Path(path)
    size = file_path.stat().st_size
    suffix = file_path.suffix.lower()
    warnings: list[str] = []
    notes: list[str] = []
    metadata: dict[str, str] = {"sha256": sha256_of(file_path)}
    title_matches: list[str] = []

    # 类型判定以文件头为准：扩展名缺失、不是已知文档扩展名、或与文件头冲突时都纠正过来，
    # 并把应当采用的扩展名交给重命名使用（端点名 downfile.jsp 不能当扩展名）。
    mime, corrected_suffix, type_note = resolve_file_type(file_path)
    if corrected_suffix:
        metadata["detected_extension"] = corrected_suffix
        metadata["canonical_extension"] = corrected_suffix
        if suffix:
            metadata["declared_extension"] = suffix
    if type_note:
        warnings.append(type_note)
    if not mime and not suffix:
        warnings.append("无法识别文件类型（文件名无扩展名且文件头未匹配）")
    # 直链守卫只管住了下载那一刻；校验环节再挡一次，避免网页被当成文件落盘
    # （实测早期版本把通知页存成了 .html 并报 success）。
    if mime.startswith("text/html"):
        warnings.append("下载到的是网页（text/html）而不是文件，不能作为目标文件")

    claimed_ok, claimed_note = check_claimed_title(claimed_title, requested_name)
    if not claimed_ok:
        warnings.append(claimed_note)

    is_pdf = mime == "application/pdf" or suffix == ".pdf" or corrected_suffix == ".pdf"
    if is_pdf:
        try:
            reader = PdfReader(str(file_path))
            page_count = len(reader.pages)
            # 名称校验保持原来的 5 页窗口，不改动既有判定。
            name_text = extract_page_text(reader, NAME_SAMPLE_PAGES)
            for token in requested_name.replace(".pdf", "").split():
                if len(token) > 1 and token in name_text:
                    title_matches.append(token)
            metadata["pages"] = str(page_count)
            matched, note = check_name_match(name_text, requested_name)
            metadata["name_match"] = note
            metadata["name_checked"] = "true"
            if not matched:
                warnings.append(note)

            # 正文实质校验：名称能对上不代表文件里真有正文。
            # 实测被它拦下的是标准平台元数据页、第三方文库付费外壳、登录墙与商业软文
            # （GB/T 8923.1-2011 那一页自己就写着"本系统暂不提供在线阅读服务"），
            # 而合法形态的网页导出（官网 HTML 正文的法律法规）因为有条款/章节结构而照常通过。
            body_text = extract_page_text(reader, BODY_SAMPLE_PAGES)
            absent = detect_absent_body(body_text, page_count)
            if absent:
                metadata["body_absent"] = "true"
                warnings.append(absent)
            else:
                scanned = detect_scanned_only(
                    page_count, body_text, count_image_pages(reader, BODY_SAMPLE_PAGES)
                )
                if scanned:
                    # 扫描件的原始内容是存在的，只是没有文本层可比对——
                    # 记进 notes（不影响 ok），并打上标记供汇总单列。
                    metadata["scan_only"] = "true"
                    notes.append(scanned)

            fragment = detect_print_fragment(name_text, page_count)
            if fragment:
                warnings.append(fragment)
            standard_number = extract_standard_number(name_text)
            if standard_number:
                metadata["standard_number"] = standard_number
        except Exception as exc:
            warnings.append(f"PDF 解析失败: {exc}")
        # 内容校验不允许静默跳过：PDF 没能解析出文本层时，无法证明它就是目标文件。
        if "name_checked" not in metadata:
            warnings.append("PDF 内容校验未执行（未能解析出文本层），不能确认内容与目标名称一致")
    elif (corrected_suffix or suffix) in {".doc", ".docx"}:
        # Word 文档以前**完全不做内容核对**（内容校验整段挂在 is_pdf 下面），于是"页面上的
        # 附件"可以被当成正文交出去。实测（2026-09-16）：请求《山东省建筑施工安全文明标准化
        # 工地管理办法》，从官方页面下到一个 19.8 KB 的 .doc，内容是"培育公示牌"附件
        # （第一条/本办法/施行 各 0 次），却因为类型识别通过被判成 `success（官方原件）`。
        kind = corrected_suffix or suffix
        try:
            word_text = extract_document_text(file_path, kind)
        except Exception as exc:
            word_text = ""
            notes.append(f"{kind} 文本抽取失败：{exc}")
        if looks_readable_text(word_text):
            matched, note = check_name_match(word_text, requested_name)
            metadata["name_match"] = note
            metadata["name_checked"] = "true"
            metadata["content_checked"] = "true"
            if not matched:
                warnings.append(note)
            number = extract_standard_number(word_text)
            if number:
                metadata["standard_number"] = number
            if kind == ".docx":
                # .docx 能完整抽到正文，跑与 PDF 同一套"里面到底有没有正文"的判据；
                # .doc（OLE2）只能粗略抽，结构判据容易误伤，所以只做名称核对。
                absent = detect_absent_body(word_text, WORD_PAGE_HINT)
                if absent:
                    metadata["body_absent"] = "true"
                    warnings.append(absent)
        else:
            # 抽不出可读文本时**不判失败**（工程上有大量正常公文抽不全），但必须如实标注：
            # 报告里能一眼看出"这一件只核对了类型与来源，内容未经核对"。
            metadata["content_checked"] = "false"
            notes.append(f"内容未核对：{kind} 抽不出可读正文，只核对了文件类型与来源")
    if not is_pdf and "name_checked" not in metadata and "content_checked" not in metadata:
        # 其它格式（.xls/.zip/.rtf/.ofd…）目前没有文本核对手段：**不判失败**，但要如实标注，
        # 让报告里能一眼看出"这一件只核对了类型与来源，内容未经核对"。
        metadata["content_checked"] = "false"
        notes.append(
            f"内容未核对：{corrected_suffix or suffix or mime or '未知类型'} 暂不支持文本核对，"
            "只核对了文件类型与来源"
        )
    if size == 0:
        warnings.append("文件大小为 0")
    return Verification(ok=size > 0 and not warnings, mime_type=mime, size=size, title_matches=title_matches, metadata=metadata, warnings=warnings, notes=notes)
